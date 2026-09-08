import streamlit as st
import pandas as pd
import json
import os
import gspread
from google.oauth2.service_account import Credentials
from docx import Document
from io import BytesIO

# ==========================================
# 1. KONFIGURASI HALAMAN & STYLING
# ==========================================
st.set_page_config(
    page_title="Sistem Terpadu Pembelajaran & Administrasi Guru",
    page_icon="🎓",
    layout="wide",
    initial_sidebar_state="expanded"
)

st.markdown("""
    <style>
    div[data-testid="stVerticalBlock"] > div[data-testid="stBlock"] {
        border-radius: 12px;
    }
    .header-box {
        background: linear-gradient(135deg, #1e3c72 0%, #2a5298 100%);
        padding: 24px;
        border-radius: 12px;
        color: white;
        margin-bottom: 25px;
        box-shadow: 0 4px 6px rgba(0,0,0,0.1);
    }
    .header-box h1 {
        color: #ffffff !important;
        margin: 0;
        font-size: 1.8rem;
    }
    .header-box p {
        color: #e0e6ed !important;
        margin: 5px 0 0 0;
        font-size: 0.95rem;
    }
    </style>
""", unsafe_allow_html=True)

# ==========================================
# 2. SINKRONISASI GOOGLE SHEETS API
# ==========================================
BASE_DIR = os.path.dirname(os.path.abspath(__file__))

@st.cache_resource
def get_gspread_client():
    scopes = [
        "https://www.googleapis.com/auth/spreadsheets",
        "https://www.googleapis.com/auth/drive"
    ]
    if "gcp_service_account" in st.secrets:
        creds = Credentials.from_service_account_info(
            st.secrets["gcp_service_account"],
            scopes=scopes
        )
        return gspread.authorize(creds)
    return None

SPREADSHEET_ID = st.secrets.get("SPREADSHEET_ID", "1o-xa_G7Vt8l4wrSrPl6AVAmhr7MXJ-AVyZf35mrW1Bc")
client = get_gspread_client()

def load_data_from_sheet(sheet_name):
    if client:
        try:
            sh = client.open_by_key(SPREADSHEET_ID)
            worksheet = sh.worksheet(sheet_name)
            data = worksheet.get_all_records()
            return pd.DataFrame(data)
        except Exception:
            return None
    return None

def save_data_to_sheet(sheet_name, df):
    if client:
        try:
            sh = client.open_by_key(SPREADSHEET_ID)
            try:
                worksheet = sh.worksheet(sheet_name)
            except Exception:
                worksheet = sh.add_worksheet(title=sheet_name, rows="100", cols="40")
            
            worksheet.clear()
            df_filled = df.fillna("")
            worksheet.update([df_filled.columns.values.tolist()] + df_filled.values.tolist())
            return True
        except Exception as e:
            st.error(f"Gagal menyimpan ke Google Sheets: {e}")
            return False
    return False

# ==========================================
# 3. DATASET SISWA DUMMY (FALLBACK)
# ==========================================
DUMMY_SISWA = {
    "Kelas 7A": [
        "Aditya Naufal Pratama", "Zidan Al Fatir", "Atika Zahara Ratifa", "Ayu Azka Fariha",
        "Dini Khoirunisa", "Enjelita Laia", "Ervan Martio Armana", "Farisman Lase", "Fauzia",
        "Hendriyanto", "M Sahel Habibillah", "Okta Dita Pranata", "Rahmad Ramadhani",
        "Yuda Agustian FitRoh", "Yulia Ramadhani", "Marveltus Hia", "Muhammad Rhaehan A"
    ],
    "Kelas 7B": [
        "Ahmad Albar", "Bagas Saputra", "Citra Kirana", "Dedi Kurniawan", "Eka Putri",
        "Fahri Hamzah", "Gita Gutawa", "Hafiz Ridho", "Indah Permata", "Joko Susilo"
    ],
    "Kelas 8": [
        "Andi Wijaya", "Budi Santoso", "Cici Paramida", "Doni Monardo", "Eva Celia"
    ],
    "Kelas 9": [
        "Abelia", "Alfinus", "Arifki", "Felicita", "Hendra", "Heri", "Hezekiel", "Lesi", 
        "Marhadi", "Melfin", "Meriati", "M. Ali", "M. Huda", "M. Alif", "M. Fauzan", 
        "Naufal", "Putra", "Putri", "Rafa", "Reno", "Rizky", "Robi", "Saputra", 
        "Selly", "Selvi", "Suci", "Yestina"
    ]
}

def load_materi_json():
    file_path = os.path.join(BASE_DIR, "materi.json")
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}

def get_data_siswa(kelas_nama):
    list_nama = DUMMY_SISWA.get(kelas_nama, DUMMY_SISWA["Kelas 7A"])
    jk_list = ["L" if i % 2 == 0 else "P" for i in range(len(list_nama))]
    return pd.DataFrame({"No": range(1, len(list_nama) + 1), "Nama": list_nama, "Jenis Kelamin": jk_list})

DATABASE_MATERI = load_materi_json()

# ==========================================
# 4. SIDEBAR PANEL
# ==========================================
with st.sidebar:
    st.markdown("### 👨‍🏫 Identitas Pengajar")
    sekolah = st.text_input("Nama Sekolah", "SMP RSUP PKB Pulau Burung")
    penyusun = st.text_input("Nama Guru / Penyusun", "Ridho Kurniawan, S.Pd.")
    
    st.divider()
    st.markdown("### 🏫 Pengaturan Kelas & Semester")
    kelas_aktif = st.selectbox(
        "Pilih Kelas Aktif",
        ["Kelas 7A", "Kelas 7B", "Kelas 8", "Kelas 9"],
        index=0
    )
    tahun = st.text_input("Tahun Pelajaran", "2026/2027")
    semester = st.selectbox("Semester", ["Ganjil", "Genap"])
    
    st.divider()
    if client:
        st.success("🟢 Otentikasi Google Sheets API Aktif!")
    else:
        st.warning("⚠️ Google Sheets API Belum Terkonfigurasi.")

DF_SISWA_AKTIF = get_data_siswa(kelas_aktif)
JUMLAH_KOLOM_NILAI = 15
KATEGORI_NILAI_OPSI = [
    "Tugas Individu", "Tugas Kelompok", "Projek / Praktik", "UTS / Mid Semester", "UAS / Akhir Semester"
]

# ==========================================
# 5. SINKRONISASI SESSION STATE & GOOGLE SHEETS
# ==========================================
key_p = f"presensi_{kelas_aktif}"
key_n = f"nilai_{kelas_aktif}"

if key_p not in st.session_state:
    df_cloud_p = load_data_from_sheet(f"Presensi_{kelas_aktif}")
    if df_cloud_p is not None and not df_cloud_p.empty:
        st.session_state[key_p] = df_cloud_p
    else:
        df_p = DF_SISWA_AKTIF.copy()
        for t in range(1, 32):
            df_p[str(t)] = ""
        st.session_state[key_p] = df_p

if key_n not in st.session_state:
    df_cloud_n = load_data_from_sheet(f"Nilai_{kelas_aktif}")
    if df_cloud_n is not None and not df_cloud_n.empty:
        st.session_state[key_n] = df_cloud_n
    else:
        df_n = DF_SISWA_AKTIF.copy()
        for i in range(1, JUMLAH_KOLOM_NILAI + 1):
            df_n[f"Nilai {i}"] = None
        st.session_state[key_n] = df_n

# ==========================================
# 6. HEADER BANNER
# ==========================================
st.markdown(f"""
    <div class="header-box">
        <h1>🎓 Portal Administrasi & Pembelajaran Guru</h1>
        <p>Selamat datang, <b>{penyusun}</b> | {sekolah} | <b>{kelas_aktif}</b> ({tahun} - Semester {semester})</p>
    </div>
""", unsafe_allow_html=True)

tab1, tab2, tab3 = st.tabs([
    "📑 Generator Modul Ajar", 
    f"📋 Buku Presensi ({kelas_aktif})", 
    f"📊 Buku Nilai & KKTP ({kelas_aktif})"
])

# ------------------------------------------
# TAB 1: GENERATOR MODUL AJAR (CUSTOMIZABLE MODEL, METODE & LKPD)
# ------------------------------------------
with tab1:
    st.subheader("⚙️ Konfigurasi Fleksibel Modul Ajar Kurikulum Merdeka")
    
    col_a, col_b = st.columns(2)
    with col_a:
        with st.container(border=True):
            st.markdown("**📖 Materi & Kurikulum**")
            if not DATABASE_MATERI:
                mapel_selected = st.selectbox("Mata Pelajaran", ["Ilmu Pengetahuan Sosial (IPS)", "Pendidikan Pancasila (PPKn)"])
                kelas_selected = st.selectbox("Jenjang Kelas Modul", ["Kelas VII", "Kelas VIII", "Kelas IX"])
                bab_selected = st.text_input("Bab / Tema Utama", "Bab 1: Kondisi Geografis dan Pelestarian SDA")
                subbab_selected = st.text_input("Sub-Materi / Subbab", "Peran Lembaga Sosial dalam Pemanfaatan SDA dan SDM")
            else:
                mapel_selected = st.selectbox("Mata Pelajaran", list(DATABASE_MATERI.keys()))
                kelas_selected = st.selectbox("Jenjang Kelas Modul", list(DATABASE_MATERI[mapel_selected].keys()))
                bab_dict = DATABASE_MATERI[mapel_selected][kelas_selected]
                bab_selected = st.selectbox("Bab / Tema Utama", list(bab_dict.keys()))
                subbab_selected = st.selectbox("Sub-Materi / Subbab", bab_dict[bab_selected])
            
            fase_selected = st.selectbox("Fase", ["Fase D", "Fase E"])
            alokasi = st.text_input("Alokasi Waktu", "2 JP (2 Pertemuan x 1 JP)")

    with col_b:
        with st.container(border=True):
            st.markdown("**🎯 Strategi & Bentuk LKPD**")
            
            model_selected = st.selectbox(
                "Model Pembelajaran",
                [
                    "Problem-Based Learning (PBL) - Deep Learning",
                    "Project-Based Learning (PjBL) - Deep Learning",
                    "Discovery / Inquiry Learning - Deep Learning",
                    "Cooperative Learning (Jigsaw/STAD) - Deep Learning"
                ]
            )
            
            metode_selected = st.selectbox(
                "Metode Pembelajaran",
                [
                    "Tanya Jawab Interaktif, Diskusi Kelompok, Presentasi & Refleksi",
                    "Studi Kasus, Bedah Masalah & Unjuk Kerja",
                    "Mind Mapping, Gallery Walk & Diskusi Pleno",
                    "Observasi Lapangan, Penyelidikan Kelompok & Presentasi"
                ]
            )
            
            tipe_lkpd = st.selectbox(
                "Pilihan Bentuk LKPD",
                [
                    "Matriks Peran & Studi Kasus Kontekstual (Tipe Review)",
                    "Analisis Kasus 3-T (Teknologi, Tantangan, Tindakan)",
                    "Mind Mapping / Peta Konsep Analitis",
                    "Rancangan Proyek Kreatif & Unjuk Karya"
                ]
            )

    def generate_custom_modul_doc(sekolah, penyusun, mapel, kelas, fase, semester, tahun, bab, subbab, alokasi, model, metode, lkpd_choice):
        doc = Document()
        
        # Title
        p_title = doc.add_paragraph()
        run_title = p_title.add_run("MODUL AJAR KURIKULUM MERDEKA\n")
        run_title.bold = True
        p_title.add_run(f"MATA PELAJARAN: {mapel.upper()}\nSTANDAR KEPUTUSAN BSKAP NOMOR 046/H/KR/2025").bold = True

        doc.add_heading("I. INFORMASI UMUM", level=1)
        
        doc.add_heading("IDENTITAS MODUL", level=2)
        p_id = doc.add_paragraph()
        p_id.add_run(f"Nama Sekolah: {sekolah}\n")
        p_id.add_run(f"Nama Penyusun: {penyusun}\n")
        p_id.add_run(f"Mata Pelajaran: {mapel}\n")
        p_id.add_run(f"Kelas / Fase / Semester: {kelas} / {fase} / {semester}\n")
        p_id.add_run(f"Bab / Tema Utama: {bab}\n")
        p_id.add_run(f"Sub-Materi Pembelajaran: {subbab}\n")
        p_id.add_run(f"Alokasi Waktu: {alokasi}\n")
        p_id.add_run(f"Tahun Pelajaran: {tahun}")

        doc.add_heading("KOMPETENSI AWAL", level=2)
        p_ka = doc.add_paragraph()
        p_ka.add_run(f"• Peserta didik telah membaca materi dasar mengenai {subbab}.\n")
        p_ka.add_run(f"• Peserta didik memiliki pengetahuan awal dalam mengidentifikasi fenomena sosial terkait {bab} di lingkungan masyarakat.")

        doc.add_heading("PROFIL PELAJAR PANCASILA", level=2)
        p_p3 = doc.add_paragraph()
        p_p3.add_run("• Beriman, Bertakwa kepada Tuhan YME, dan Berakhlak Mulia: Menumbuhkan kesadaran moral & etika sosial.\n")
        p_p3.add_run("• Bernalar Kritis: Mampu menganalisis masalah dan fenomena secara objektif.\n")
        p_p3.add_run("• Gotong Royong: Berkolaborasi dan berdiskusi secara efektif dalam kelompok.\n")
        p_p3.add_run(f"• Kreatif: Menghasilkan gagasan/solusi inovatif terkait {subbab}.")

        doc.add_heading("SARANA DAN PRASARANA", level=2)
        p_sp = doc.add_paragraph()
        p_sp.add_run("• Media Pembelajaran: Papan Tulis, Slide Presentasi, Kartu Kasus/Artikel, Lembar Kerja Peserta Didik (LKPD).\n")
        p_sp.add_run(f"• Sumber Belajar Utama: Buku Paket / LKS Siswa {mapel}, Lingkungan Sekitar Sekolah.")

        doc.add_heading("TARGET PESERTA DIDIK & STRATEGI PEMBELAJARAN", level=2)
        p_tp_d = doc.add_paragraph()
        p_tp_d.add_run("• Target Peserta Didik: Peserta didik reguler / tipikal.\n")
        p_tp_d.add_run(f"• Model Pembelajaran: {model}\n")
        p_tp_d.add_run(f"• Metode Pembelajaran: {metode}")

        doc.add_heading("II. KOMPONEN INTI", level=1)
        
        doc.add_heading("TUJUAN PEMBELAJARAN (TP)", level=2)
        p_tp = doc.add_paragraph()
        p_tp.add_run(f"1. Peserta didik mampu menjelaskan serta menganalisis konsep {subbab} secara tepat.\n")
        p_tp.add_run(f"2. Peserta didik mampu mengidentifikasi serta memecahkan studi kasus kontekstual yang berkaitan dengan {bab}.\n")
        p_tp.add_run(f"3. Peserta didik mampu menyajikan hasil diskusi dan analisis kelompok secara komunikatif di depan kelas.")

        doc.add_heading("PEMAHAMAN BERMAKNA (MEANINGFUL LEARNING)", level=2)
        p_mb = doc.add_paragraph()
        p_mb.add_run(f"Pemahaman terhadap {subbab} membantu peserta didik menyadari peran aktifnya sebagai warga negara yang bijak, kritis, dan bertanggung jawab di tengah kehidupan masyarakat.")

        doc.add_heading("PERTANYAAN PEMANTIK", level=2)
        p_pp = doc.add_paragraph()
        p_pp.add_run(f"1. Mengapa topik {subbab} sangat penting dalam kehidupan kita sehari-hari?\n")
        p_pp.add_run(f"2. Dampak apa yang akan terjadi jika prinsip {bab} tidak diterapkan dengan baik di masyarakat?")

        doc.add_heading("III. KEGIATAN PEMBELAJARAN DETAIL", level=1)
        
        doc.add_heading("PENDAHULUAN (15 MENIT) - Mindful Start", level=2)
        p_pen = doc.add_paragraph()
        p_pen.add_run("• Orientasi Kelas: Guru menyapa peserta didik, memimpin doa, dan mengecek presensi kehadiran.\n")
        p_pen.add_run(f"• Apersepsi Mindful: Guru memberikan pemantik singkat untuk memancing pengetahuan awal siswa terkait {subbab}.\n")
        p_pen.add_run(f"• Penyampaian Tujuan: Guru menjelaskan tujuan pembelajaran dan skenario kegiatan dengan model {model}.")

        doc.add_heading(f"KEGIATAN INTI (50 MENIT) - Meaningful & Joyful ({model})", level=2)
        p_inti = doc.add_paragraph()
        p_inti.add_run(f"1. Orientasi / Stimulasi Masalah (~10 Menit):\n   Guru menyajikan contoh fenomena/masalah nyata di masyarakat terkait {subbab}.\n\n")
        p_inti.add_run(f"2. Mengorganisasikan Kelompok (~5 Menit):\n   Siswa dibagi menjadi kelompok kecil (3-4 siswa) dan dibagikan {lkpd_choice}.\n\n")
        p_inti.add_run(f"3. Penyelidikan & Penerapan Metode ({metode}) (~20 Menit):\n   Siswa berdiskusi menyelesaikan tugas pada LKPD dengan bimbingan terarah (scaffolding) dari guru.\n\n")
        p_inti.add_run(f"4. Mengembangkan & Menyajikan Hasil Karya (~15 Menit):\n   Perwakilan kelompok mempresentasikan hasil analisis LKPD di depan kelas dan ditanggapi kelompok lain.")

        doc.add_heading("PENUTUP (15 MENIT) - Deep Reflection", level=2)
        p_penutup = doc.add_paragraph()
        p_penutup.add_run(f"• Sintesis & Rangkuman: Guru bersama peserta didik merangkum poin kunci dari {subbab}.\n")
        p_penutup.add_run("• Refleksi Pembelajaran: Peserta didik menuliskan 1 kalimat refleksi: 'Satu hal paling penting yang saya pelajari hari ini adalah ...'\n")
        p_penutup.add_run("• Apresiasi & Tindak Lanjut: Guru memberikan pujian atas partisipasi aktif siswa dan menutup kelas dengan doa bersama.")

        doc.add_heading("IV. ASESMEN PEMBELAJARAN (PENILAIAN)", level=1)
        p_as = doc.add_paragraph()
        p_as.add_run("• Asesmen Sikap: Observasi Profil Pelajar Pancasila (Bernalar Kritis, Gotong Royong, Mandiri).\n")
        p_as.add_run(f"• Asesmen Formatif: Penilaian Diskusi Kelompok & Kinerja Pengerjaan {lkpd_choice}.\n")
        p_as.add_run("• Asesmen Sumatif: Hasil Akhir Pengerjaan Lembar Kerja Peserta Didik (LKPD).")

        doc.add_heading("V. LAMPIRAN: LEMBAR KERJA PESERTA DIDIK (LKPD)", level=1)
        
        doc.add_heading(f"LEMBAR KERJA PESERTA DIDIK ({lkpd_choice.upper()})", level=2)
        p_lkpd = doc.add_paragraph()
        p_lkpd.add_run(f"Nama Kelompok : ....................................\nKelas : {kelas}\nAnggota Kelompok : 1. ..... 2. ..... 3. ..... 4. .....\n\n")
        
        # MENYESUAIKAN ISI LKPD SESUAI PILIHAN GURU
        if "Matriks Peran" in lkpd_choice:
            p_lkpd.add_run("BAGIAN A: MATRIKS ANALISIS TABEL\nIsilah tabel di bawah ini berdasarkan bacaan/materi yang telah dipelajari!\n").bold = True
            table_m = doc.add_table(rows=4, cols=3)
            table_m.style = 'Table Grid'
            table_m.rows[0].cells[0].text = 'No'
            table_m.rows[0].cells[1].text = 'Kategori / Komponen'
            table_m.rows[0].cells[2].text = f'Analisis Peran dalam {subbab}'
            for i in range(1, 4):
                table_m.rows[i].cells[0].text = str(i)
            
            doc.add_paragraph("\nBAGIAN B: STUDI KASUS KONTEKSTUAL\nBacalah kasus di bawah ini lalu berikan solusinya:\n\"Di sebuah daerah terjadi permasalahan sosial dan lingkungan terkait pemanfaatan sumber daya. Jelaskan lembaga/pihak mana saja yang harus memberikan solusi!\"\nJawaban: ....................................................................................................................................................................................................................................................................................").bold = True
            
        elif "3-T" in lkpd_choice:
            p_lkpd.add_run("ANALISIS KASUS BERDASARKAN RUMUS 3-T:\n").bold = True
            p_lkpd.add_run("1. TEKNOLOGI / PENDORONG (T-1):\n   Sarana/faktor apa yang menjadi pendorong utama munculnya fenomena ini?\n   Jawaban: ....................................................................................................\n\n")
            p_lkpd.add_run("2. TANTANGAN DAN DAMPAK (T-2):\n   Tuliskan 2 dampak positif dan 2 dampak negatif dari fenomena ini!\n   - Positif: 1. ..... 2. .....\n   - Negatif: 1. ..... 2. .....\n\n")
            p_lkpd.add_run("3. TINDAKAN DAN SOLUSI (T-3):\n   Tindakan nyata apa yang harus dilakukan generasi muda untuk mengatasi tantangan tersebut?\n   Jawaban: ....................................................................................................\n")
            
        elif "Mind Mapping" in lkpd_choice:
            p_lkpd.add_run("PETUNJUK PETA KONSEP (MIND MAPPING):\n").bold = True
            p_lkpd.add_run(f"1. Buatlah Peta Konsep / Mind Map di lembar ini mengenai hubungan antara {bab} dengan {subbab}!\n")
            p_lkpd.add_run("2. Gunakan kata kunci, cabang utama, cabang pembantu, serta warna/gambar menarik.\n")
            p_lkpd.add_run("3. Jelaskan secara singkat alur peta konsep kelompokmu di depan kelas!\n\n[ KOTAK LEMBAR KERJA PETA KONSEP ]\n\n\n\n\n")
            
        else: # Proyek Kreatif
            p_lkpd.add_run("PERANCANGAN PROYEK KREATIF KELOMPOK:\n").bold = True
            p_lkpd.add_run(f"1. Judul Produk / Karya: (Poster / Infografis / Video Short mengenai {subbab})\n")
            p_lkpd.add_run("2. Alasan Pemilihan Produk: ....................................................................................\n")
            p_lkpd.add_run("3. Langkah-Langkah Pembuatan Proyek:\n   a. .....\n   b. .....\n   c. .....\n")
            p_lkpd.add_run("4. Pembagian Tugas Anggota Kelompok: ....................................................................................\n")

        # Rubrik Penilaian
        doc.add_heading("RUBRIK PENILAIAN KELOMPOK", level=2)
        table_r = doc.add_table(rows=3, cols=5)
        table_r.style = 'Table Grid'
        hdr_cells = table_r.rows[0].cells
        hdr_cells[0].text = 'Kriteria'
        hdr_cells[1].text = 'Sangat Baik (4)'
        hdr_cells[2].text = 'Baik (3)'
        hdr_cells[3].text = 'Cukup (2)'
        hdr_cells[4].text = 'Perlu Bimbingan (1)'
        
        row1 = table_r.rows[1].cells
        row1[0].text = 'Gotong Royong'
        row1[1].text = 'Semua aktif & kompak'
        row1[2].text = 'Sebagian besar aktif'
        row1[3].text = 'Hanya 1-2 siswa dominan'
        row1[4].text = 'Pasif'

        row2 = table_r.rows[2].cells
        row2[0].text = 'Bernalar Kritis'
        row2[1].text = 'Analisis sangat dalam & solutif'
        row2[2].text = 'Analisis cukup mendalam'
        row2[3].text = 'Jawaban umum'
        row2[4].text = 'Belum mampu menjawab'

        # Tanda Tangan
        doc.add_paragraph("\n\n")
        p_ttd = doc.add_paragraph()
        p_ttd.add_run("Mengetahui,\t\t\t\t\t\tGuru Mata Pelajaran\nKepala Sekolah\n\n\n\n").bold = True
        p_ttd.add_run(f"( .................................................... )\t\t\t\t({penyusun})")

        buf = BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf

    buf_doc = generate_custom_modul_doc(
        sekolah, penyusun, mapel_selected, kelas_selected, fase_selected, 
        semester, tahun, bab_selected, subbab_selected, alokasi, 
        model_selected, metode_selected, tipe_lkpd
    )

    st.download_button(
        label="📥 Download Modul Ajar Custom (.docx)",
        data=buf_doc,
        file_name=f"Modul_Ajar_{mapel_selected}_{kelas_selected}.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        type="primary"
    )

# ------------------------------------------
# TAB 2: BUKU PRESENSI
# ------------------------------------------
with tab2:
    st.subheader(f"📖 Buku Presensi Harian ({kelas_aktif})")
    bulan_presensi = st.selectbox(
        "Pilih Bulan Presensi",
        ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober", "November", "Desember"],
        index=8,
        key=f"bln_{kelas_aktif}"
    )

    df_p_curr = st.session_state[key_p].copy()
    tgl_cols = [str(t) for t in range(1, 32)]

    col_config_p = {
        "No": st.column_config.NumberColumn("No", disabled=True, width="small"),
        "Nama": st.column_config.TextColumn("Nama Siswa", disabled=True, width="medium"),
        "Jenis Kelamin": st.column_config.TextColumn("JK", disabled=True, width="small"),
    }
    for t in tgl_cols:
        col_config_p[t] = st.column_config.TextColumn(t, width="small")

    df_upper = df_p_curr[tgl_cols].fillna("").apply(lambda x: x.astype(str).str.upper())
    df_p_curr["H"] = (df_upper == "H").sum(axis=1)
    df_p_curr["S"] = (df_upper == "S").sum(axis=1)
    df_p_curr["I"] = (df_upper == "I").sum(axis=1)
    df_p_curr["A"] = (df_upper == "A").sum(axis=1)

    col_config_p["H"] = st.column_config.NumberColumn("H", disabled=True, width="small")
    col_config_p["S"] = st.column_config.NumberColumn("S", disabled=True, width="small")
    col_config_p["I"] = st.column_config.NumberColumn("I", disabled=True, width="small")
    col_config_p["A"] = st.column_config.NumberColumn("A", disabled=True, width="small")

    edited_p = st.data_editor(
        df_p_curr,
        column_config=col_config_p,
        disabled=["No", "Nama", "Jenis Kelamin", "H", "S", "I", "A"],
        hide_index=True,
        use_container_width=True,
        key=f"editor_p_{kelas_aktif}"
    )
    
    st.session_state[key_p] = edited_p[DF_SISWA_AKTIF.columns.tolist() + tgl_cols]

    if st.button(f"💾 Simpan Presensi {kelas_aktif} ke Google Sheets", type="primary"):
        if save_data_to_sheet(f"Presensi_{kelas_aktif}", edited_p):
            st.success("✅ Data Presensi Berhasil Disimpan Permanen ke Google Sheets!")

# ------------------------------------------
# TAB 3: BUKU NILAI & KKTP
# ------------------------------------------
with tab3:
    st.subheader(f"📊 Buku Nilai & Akumulasi Realtime ({kelas_aktif})")
    
    with st.expander("📌 Pengaturan Kategori/Jenis Penilaian per Kolom", expanded=True):
        st.write("Silakan tentukan jenis penilaian untuk masing-masing kolom di bawah ini:")
        cols_kat = st.columns(5)
        kategori_terpilih = {}
        for i in range(1, JUMLAH_KOLOM_NILAI + 1):
            col_idx = (i - 1) % 5
            with cols_kat[col_idx]:
                kategori_terpilih[f"Nilai {i}"] = st.selectbox(
                    f"Kolom Nilai {i}",
                    options=KATEGORI_NILAI_OPSI,
                    index=0 if i <= 10 else 1,
                    key=f"kat_n_{kelas_aktif}_{i}"
                )

    df_nilai_current = st.session_state[key_n].copy()
    
    column_config_n = {
        "No": st.column_config.NumberColumn("No", disabled=True, width="small"),
        "Nama": st.column_config.TextColumn("Nama Siswa", disabled=True, width="medium"),
        "Jenis Kelamin": st.column_config.TextColumn("JK", disabled=True, width="small"),
    }
    
    kolom_nilai_keys = [f"Nilai {i}" for i in range(1, JUMLAH_KOLOM_NILAI + 1)]
    
    for k in kolom_nilai_keys:
        label_kat = kategori_terpilih[k]
        column_config_n[k] = st.column_config.NumberColumn(
            f"{k} ({label_kat})",
            min_value=0.0,
            max_value=100.0,
            format="%.1f",
            width="medium"
        )
    
    df_numeric = df_nilai_current[kolom_nilai_keys].apply(pd.to_numeric, errors='coerce')
    df_nilai_current["📊 Nilai Akhir"] = df_numeric.mean(axis=1).round(2)
    
    column_config_n["📊 Nilai Akhir"] = st.column_config.NumberColumn(
        "📊 Nilai Akhir", disabled=True, format="%.2f", width="medium"
    )

    edited_n = st.data_editor(
        df_nilai_current,
        column_config=column_config_n,
        disabled=["No", "Nama", "Jenis Kelamin", "📊 Nilai Akhir"],
        hide_index=True,
        use_container_width=True,
        key=f"editor_n_{kelas_aktif}"
    )
    
    st.session_state[key_n] = edited_n[DF_SISWA_AKTIF.columns.tolist() + kolom_nilai_keys]

    st.markdown("---")
    if st.button(f"💾 Simpan Buku Nilai {kelas_aktif} ke Google Sheets", type="primary"):
        if save_data_to_sheet(f"Nilai_{kelas_aktif}", edited_n):
            st.success("✅ Data Buku Nilai Berhasil Disimpan Permanen ke Google Sheets!")
